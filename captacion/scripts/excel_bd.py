import json
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
keep = json.load(open('bd_icp_keep.json')); out = json.load(open('bd_icp_out.json'))
wb = Workbook(); H = Font(bold=True, color='FFFFFF'); F = PatternFill('solid', fgColor='1F3A5F')
def hdr(ws, cols, widths):
    ws.append(cols)
    for i,c in enumerate(cols,1):
        cell = ws.cell(1,i); cell.font=H; cell.fill=F
        ws.column_dimensions[get_column_letter(i)].width = widths[i-1]
    ws.freeze_panes = 'A2'
ws = wb.active; ws.title = 'Resumen'
ws.append(['Base ICP del plan trimestral · empresas que ya invierten en captación · 15-oct-2026']); ws['A1'].font = Font(bold=True, size=13)
ws.append(['Fuente: Apollo, búsqueda de empresas (gratis, sin créditos). España, 6-200 empleados. Señales: anuncios activos (píxel Meta / conversión Google), CRM instalado (HubSpot, Salesforce, Pipedrive, Zoho, ActiveCampaign), equipo comercial (2+ personas en ventas).'])
ws.append(['Niveles del plan: A = 1ª línea (formación, reformas e instalaciones, clínicas) con anuncios activos y además CRM o equipo comercial. B = 1ª o 2ª línea (inmobiliaria, servicios B2B, solar) con al menos una señal. C = resto, solo exploración. La economía (facturación de 20-30k €/mes, inversión de 1.000 €/mes) queda «pendiente de validar» en todas: Apollo no la confirma y no se confunde falta de dato con falta de encaje.'])
ws.append(['Aún no hay personas ni correos: esto es la lista de cuentas. La tipología se asigna automáticamente por nombre y búsqueda, así que revisar antes de cargar.'])
ws.append([])
ws.append(['Tipología','Total','A','B','C','Con anuncios','Con CRM','Con equipo comercial'])
for i in range(1,9): ws.cell(6,i).font=H; ws.cell(6,i).fill=F
t = Counter(a['tipologia'] for a in keep)
for k,v in t.most_common():
    sub=[a for a in keep if a['tipologia']==k]; p=Counter(a['prioridad'] for a in sub)
    ws.append([k,v,p['A'],p['B'],p['C'],sum('anuncios' in a['senales'] for a in sub),sum('CRM' in a['senales'] for a in sub),sum('equipo comercial' in a['senales'] for a in sub)])
p=Counter(a['prioridad'] for a in keep)
ws.append(['TOTAL',len(keep),p['A'],p['B'],p['C'],sum('anuncios' in a['senales'] for a in keep),sum('CRM' in a['senales'] for a in keep),sum('equipo comercial' in a['senales'] for a in keep)])
ws.cell(ws.max_row,1).font=Font(bold=True)
ws.append([]); ws.append(['Descartadas','Nº'])
for m,n in Counter(a['motivo'] for a in out).most_common(): ws.append([m,n])
ws.column_dimensions['A'].width=40
for c in 'BCDEFGH': ws.column_dimensions[c].width=14
b = wb.create_sheet('Base')
hdr(b, ['ICP','Línea','Tipología','Empresa','Dominio','Web','Señales','Nº señales','Economía','Apollo org id','Estado'], [6,13,28,45,30,38,38,11,20,26,14])
for a in keep:
    b.append([a['prioridad'],a['linea'],a['tipologia'],a['nombre'],a['dominio'],a['web'],' + '.join(a['senales']),a['n_senales'],a['economia'],a['id'],'sin contactar'])
b.auto_filter.ref = b.dimensions
d = wb.create_sheet('Descartadas')
hdr(d, ['Motivo','Tipología','Empresa','Dominio','Señales'], [40,28,45,30,38])
for a in sorted(out, key=lambda a:(a['motivo'],a['nombre'].lower())):
    d.append([a['motivo'],a['tipologia'],a['nombre'],a['dominio'],' + '.join(a['senales'])])
d.auto_filter.ref = d.dimensions
wb.save('/home/user/qualivo/captacion/datos/base-icp-q4-15oct.xlsx')
import csv
with open('/home/user/qualivo/captacion/datos/base-icp-q4-15oct.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['icp','linea','tipologia','empresa','dominio','web','senales','n_senales','economia','apollo_org_id'])
    for a in keep: w.writerow([a['prioridad'],a['linea'],a['tipologia'],a['nombre'],a['dominio'],a['web'],' + '.join(a['senales']),a['n_senales'],a['economia'],a['id']])
print('ok')
