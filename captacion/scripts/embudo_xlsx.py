#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Genera el workbook con diseño del embudo Qualivo a partir de los CSV de
# captacion/datos/ (funnel-diario, campanas-diario, embudo-ghl).
# Uso: python3 embudo_xlsx.py [YYYY-MM-DD] [salida.xlsx]
# Pestañas: Resumen (hoy/semana/mes) · Campañas · Pipeline · Histórico.
import csv, sys, os, datetime, zoneinfo, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TZ = zoneinfo.ZoneInfo("Europe/Madrid")
HOY = sys.argv[1] if len(sys.argv) > 1 else datetime.datetime.now(TZ).date().isoformat()
RAIZ = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")
SALIDA = sys.argv[2] if len(sys.argv) > 2 else os.path.join(RAIZ, "captacion", "datos", f"embudo-{HOY}.xlsx")
D = lambda f: list(csv.DictReader(open(os.path.join(RAIZ, "captacion", "datos", f))))

NAVY = "1F3757"; ACENTO = "2E86AB"; GRIS = "F2F4F7"; VERDE = "1E7A46"; ROJO = "B3261E"
F = lambda **k: Font(name="Arial", **k)
TH = Font(name="Arial", bold=True, color="FFFFFF", size=10)
FILL_H = PatternFill("solid", fgColor=NAVY)
FILL_S = PatternFill("solid", fgColor=GRIS)
BORDE = Border(bottom=Side(style="thin", color="D9D9D9"))
CENT = Alignment(horizontal="center")

def cabecera(ws, fila, cols, anchos):
    for i, (c, a) in enumerate(zip(cols, anchos), 1):
        cel = ws.cell(row=fila, column=i, value=c)
        cel.font = TH; cel.fill = FILL_H; cel.alignment = CENT
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)

def titulo(ws, texto, sub=""):
    ws["A1"] = texto; ws["A1"].font = F(bold=True, size=15, color=NAVY)
    if sub:
        ws["A2"] = sub; ws["A2"].font = F(size=10, color="777777")

wb = Workbook(); wb.remove(wb.active)

# --- Resumen -------------------------------------------------------------------
funnel = D("funnel-diario.csv")
lunes = (datetime.date.fromisoformat(HOY) - datetime.timedelta(days=datetime.date.fromisoformat(HOY).weekday())).isoformat()
mes = HOY[:8] + "01"
CANALES = ["email-frio", "whatsapp", "voz", "linkedin", "radiografia", "web-blog"]
ETIQ = {"email-frio": "Email frío", "whatsapp": "WhatsApp", "voz": "Voz (Raquel)",
        "linkedin": "LinkedIn", "radiografia": "Radiografía", "web-blog": "Web / Blog"}
METS = ["volumen", "aperturas", "clics", "respuestas", "conversaciones", "reuniones"]

def suma(desde):
    t = collections.defaultdict(collections.Counter)
    for r in funnel:
        if desde <= r["fecha"] <= HOY:
            for m in METS:
                if str(r.get(m) or "").isdigit(): t[r["canal"]][m] += int(r[m])
    return t

ws = wb.create_sheet("Resumen")
titulo(ws, "EMBUDO QUALIVO", f"Actualizado {HOY} · la métrica que manda: reuniones cualificadas")
fila = 4
for nombre, desde in (("HOY", HOY), (f"SEMANA (desde {lunes})", lunes), (f"MES (desde {mes})", mes)):
    ws.cell(row=fila, column=1, value=nombre).font = F(bold=True, size=11, color=ACENTO)
    fila += 1
    cabecera(ws, fila, ["Canal", "Volumen", "Aperturas", "Clics", "Respuestas", "Conversaciones", "Reuniones"],
             [16, 10, 11, 8, 12, 15, 11])
    t = suma(desde)
    for j, canal in enumerate(CANALES):
        fila += 1
        ws.cell(row=fila, column=1, value=ETIQ[canal]).font = F(size=10)
        for i, m in enumerate(METS, 2):
            c = ws.cell(row=fila, column=i, value=t.get(canal, {}).get(m, 0))
            c.font = F(size=10); c.alignment = CENT; c.border = BORDE
            if j % 2: c.fill = FILL_S
        if j % 2: ws.cell(row=fila, column=1).fill = FILL_S
        ws.cell(row=fila, column=1).border = BORDE
    fila += 3
ws.freeze_panes = "A3"

# --- Campañas ------------------------------------------------------------------
ws = wb.create_sheet("Campañas")
titulo(ws, "CAMPAÑAS DE EMAIL", f"Día {HOY} · depósito = leads cargados sin empezar")
cabecera(ws, 4, ["Campaña", "Estado", "Enviados", "Aperturas", "Clics", "Respuestas", "Rebotes", "Depósito"],
         [46, 10, 10, 11, 8, 12, 9, 10])
fila = 4
camps = [r for r in D("campanas-diario.csv") if r["fecha"] == HOY]
camps.sort(key=lambda r: -int(r["enviados"] or 0))
for j, r in enumerate(camps):
    fila += 1
    vals = [r["campana"], r["estado"], int(r["enviados"] or 0), int(r["aperturas"] or 0),
            int(r["clics"] or 0), int(r["respuestas"] or 0), int(r["rebotes"] or 0),
            int(r["deposito_sin_empezar"]) if str(r["deposito_sin_empezar"]).isdigit() else ""]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=fila, column=i, value=v)
        c.font = F(size=10); c.border = BORDE
        if i > 1: c.alignment = CENT
        if j % 2: c.fill = FILL_S
    if int(r["respuestas"] or 0) > 0:
        ws.cell(row=fila, column=6).font = F(size=10, bold=True, color=VERDE)
    dep = r["deposito_sin_empezar"]
    if str(dep).isdigit() and int(dep) == 0 and r["estado"] == "ACTIVE":
        ws.cell(row=fila, column=8).font = F(size=10, bold=True, color=ROJO)
fila += 1
ws.cell(row=fila, column=1, value="TOTAL").font = F(bold=True, size=10)
for i, col in enumerate("CDEFG", 3):
    ws.cell(row=fila, column=i, value=f"=SUM({col}5:{col}{fila-1})").font = F(bold=True, size=10)
    ws.cell(row=fila, column=i).alignment = CENT

# --- Pipeline ------------------------------------------------------------------
ws = wb.create_sheet("Pipeline")
titulo(ws, "PIPELINE GHL", f"Foto del {HOY} · dónde está cada oportunidad")
cabecera(ws, 4, ["Pipeline", "Etapa", "Oportunidades", "Valor €"], [20, 30, 14, 12])
fila = 4
for j, r in enumerate([r for r in D("embudo-ghl.csv") if r["fecha"] == HOY]):
    fila += 1
    for i, v in enumerate([r["pipeline"], r["etapa"], int(r["oportunidades"]),
                           int(float(r["valor_eur"] or 0))], 1):
        c = ws.cell(row=fila, column=i, value=v)
        c.font = F(size=10); c.border = BORDE
        if i >= 3: c.alignment = CENT
        if i == 4: c.number_format = "#,##0"
        if j % 2: c.fill = FILL_S
    if "Reunión" in r["etapa"] or "Agendada" in r["etapa"]:
        ws.cell(row=fila, column=2).font = F(size=10, bold=True, color=VERDE)

# --- Histórico -----------------------------------------------------------------
ws = wb.create_sheet("Histórico")
titulo(ws, "HISTÓRICO DIARIO POR CANAL", "Una fila por día y canal · fuente: captacion/datos/funnel-diario.csv")
cabecera(ws, 4, ["Fecha", "Canal", "Volumen", "Aperturas", "Clics", "Respuestas", "Conversaciones", "Reuniones", "Notas"],
         [11, 13, 9, 10, 7, 11, 14, 10, 44])
fila = 4
for j, r in enumerate(funnel):
    fila += 1
    for i, k in enumerate(["fecha", "canal", "volumen", "aperturas", "clics", "respuestas",
                           "conversaciones", "reuniones", "notas"], 1):
        v = r.get(k, "")
        c = ws.cell(row=fila, column=i, value=int(v) if str(v).isdigit() else v)
        c.font = F(size=9); c.border = BORDE
        if 3 <= i <= 8: c.alignment = CENT
        if j % 2: c.fill = FILL_S

wb.save(SALIDA)
print(SALIDA)
